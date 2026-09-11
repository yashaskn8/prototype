from io import BytesIO
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone

from core.models import ServiceCall, Customer, Site, Asset, Project, StaffProfile

EXPORT_HEADERS = [
    "Servy ID", "Call Type", "Complaint Type", "Complaint", "Created", "Updated",
    "Customer", "Site", "Project", "Asset", "Asset Code", "Technician", "Status", "Priority",
    "Contact Name", "Contact Phone", "Contact Email", "Resolution"
]


def _safe_excel(value):
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def export_calls_xlsx(tenant):
    wb = Workbook()
    ws = wb.active
    ws.title = "Call Register"
    ws.append(EXPORT_HEADERS)
    qs = ServiceCall.objects.filter(tenant=tenant).select_related("customer", "site", "project", "asset", "technician")
    for call in qs.iterator(chunk_size=500):
        ws.append([
            call.servy_id,
            _safe_excel(call.call_type),
            _safe_excel(call.complaint_type),
            _safe_excel(call.complaint_text),
            call.created_at.replace(tzinfo=None) if call.created_at else "",
            call.updated_at.replace(tzinfo=None) if call.updated_at else "",
            _safe_excel(call.customer.name),
            _safe_excel(call.site.name if call.site else ""),
            _safe_excel(call.project.code if call.project else ""),
            _safe_excel(call.asset.name if call.asset else ""),
            _safe_excel(call.asset.asset_code if call.asset else ""),
            _safe_excel(call.technician.full_name if call.technician else ""),
            call.get_status_display(), call.get_priority_display(),
            _safe_excel(call.contact_name), _safe_excel(call.contact_phone), _safe_excel(call.contact_email),
            _safe_excel(call.resolution_text),
        ])
    ws.freeze_panes = "A2"
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 12), 42)
    stream = BytesIO()
    wb.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="call_register_{tenant.slug}.xlsx"'
    response["X-Content-Type-Options"] = "nosniff"
    return response


def import_calls_xlsx(tenant, uploaded_file):
    if Path(uploaded_file.name).suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx files can be imported.")
    max_rows = settings.SERVY_MAX_IMPORT_ROWS
    # XLSX is a ZIP archive. Bound its expanded size before openpyxl touches it
    # to reduce zip-bomb / memory-exhaustion risk.
    try:
        uploaded_file.seek(0)
        with zipfile.ZipFile(uploaded_file) as zf:
            infos = zf.infolist()
            total_uncompressed = sum(i.file_size for i in infos)
            if total_uncompressed > settings.SERVY_MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError("The Excel workbook expands beyond the configured safety limit.")
            if len(infos) > 500:
                raise ValueError("The Excel workbook contains too many internal files.")
        uploaded_file.seek(0)
        wb = load_workbook(uploaded_file, data_only=True, read_only=True, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The uploaded Excel workbook could not be read safely.") from exc
    if len(wb.sheetnames) > 20:
        raise ValueError("The Excel workbook contains too many sheets.")
    ws = wb.active
    if ws.max_column > 80:
        raise ValueError("The Excel worksheet contains too many columns.")
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {h.lower(): i for i, h in enumerate(headers)}

    def val(row, *names):
        for name in names:
            idx = index.get(name.lower())
            if idx is not None and idx < len(row):
                return row[idx].value
        return None

    created = 0
    skipped = 0
    for row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if row_number > max_rows + 1:
            raise ValueError(f"Import stopped: workbook exceeds the {max_rows} row safety limit.")
        customer_name = val(row, "Customer")
        complaint = val(row, "Complaint", "Complaint Type")
        if not customer_name or not complaint:
            skipped += 1
            continue
        customer, _ = Customer.objects.get_or_create(tenant=tenant, name=str(customer_name).strip()[:160])
        site_name = val(row, "Site")
        site = None
        if site_name:
            site, _ = Site.objects.get_or_create(tenant=tenant, customer=customer, name=str(site_name).strip()[:150])

        asset = None
        asset_code = val(row, "Asset Code")
        asset_name = val(row, "Asset")
        if asset_code:
            asset = Asset.objects.filter(tenant=tenant, customer=customer, asset_code=str(asset_code).strip()).first()
        if not asset and asset_name:
            asset = Asset.objects.filter(tenant=tenant, customer=customer, name=str(asset_name).strip()).first()

        project = None
        project_code = val(row, "Project")
        if project_code:
            project = Project.objects.filter(tenant=tenant, customer=customer, code=str(project_code).strip()).first()

        technician = None
        technician_name = val(row, "Technician")
        if technician_name:
            technician = StaffProfile.objects.filter(tenant=tenant, full_name=str(technician_name).strip(), is_active=True).first()

        proposed_id = val(row, "Servy ID")
        try:
            proposed_id = int(proposed_id)
        except Exception:
            proposed_id = (ServiceCall.objects.filter(tenant=tenant).order_by("-servy_id").values_list("servy_id", flat=True).first() or 42830) + 1
        if ServiceCall.objects.filter(tenant=tenant, servy_id=proposed_id).exists():
            skipped += 1
            continue

        status = str(val(row, "Status") or "open").strip().lower().replace(" ", "_")
        valid_status = {k for k, _ in ServiceCall.STATUS_CHOICES}
        if status not in valid_status:
            status = "open"
        priority = str(val(row, "Priority") or "normal").strip().lower()
        valid_priority = {k for k, _ in ServiceCall.PRIORITIES}
        if priority not in valid_priority:
            priority = "normal"

        ServiceCall.objects.create(
            tenant=tenant,
            servy_id=proposed_id,
            call_type=str(val(row, "Call Type") or "Service")[:80],
            complaint_type=str(val(row, "Complaint Type") or "Imported Complaint")[:180],
            complaint_text=str(val(row, "Complaint") or "")[:10000],
            customer=customer,
            site=site,
            project=project,
            asset=asset,
            technician=technician,
            status=status,
            priority=priority,
            contact_name=str(val(row, "Contact Name") or customer.contact_name)[:150],
            contact_phone=str(val(row, "Contact Phone") or customer.phone)[:30],
            contact_email=str(val(row, "Contact Email") or customer.email)[:254],
            resolution_text=str(val(row, "Resolution") or "")[:10000],
            created_at=timezone.now(),
        )
        created += 1
    return created, skipped
