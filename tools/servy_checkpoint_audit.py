"""Read-only Servy checkpoint comparison on one frozen retrieval candidate set.

RESEARCH-ONLY: This tool does not train, write a database, change indexes, or
promote models. Reused weakly supervised dev labels are not independent accuracy
evidence. Private cases and checkpoints stay outside Git.

Example (Colab; run from repository root):
  python tools/servy_checkpoint_audit.py \
    --cases /content/servy_colab_dev.jsonl \
    --corpus /content/servy_colab_corpus.jsonl \
    --model pretrained=BAAI/bge-reranker-v2-m3 \
    --model v2=/content/servy_reranker_provisional_v2 \
    --model v3_stage1=/content/servy_reranker_v3_stage1 \
    --model v3_final=/content/servy_reranker_v3_research \
    --out-dir /content/servy_checkpoint_audit

Frozen --candidates JSONL rows: {"query_id":"...", "chunk_ids":["...", ...]}.
If omitted, Qwen3 generates one tenant-filtered candidate ranking per case.
Tenant filtering alone does NOT reproduce Servy's full production authorization.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path


def load_jsonl(path):
    rows = []
    with Path(path).open(encoding="utf-8") as f:
        for line_number, line in enumerate(f, 1):
            if line.strip():
                try:
                    rows.append(json.loads(line))
                except json.JSONDecodeError as exc:
                    raise ValueError(f"{path}:{line_number}: invalid JSON") from exc
    if not rows:
        raise ValueError(f"{path}: no rows")
    return rows


def _id(value):
    if value is None or str(value).strip() == "":
        raise ValueError("Missing required ID")
    return str(value).strip()


def _field(row, *names):
    for name in names:
        if row.get(name) is not None and str(row[name]).strip():
            return row[name]
    raise ValueError(f"Missing required field (one of {names})")


def normalize_inputs(cases, corpus):
    """Fail closed on missing metadata, ID collisions and absent gold chunks."""
    chunks = {}
    for row in corpus:
        cid = _id(_field(row, "chunk_id", "passage_id"))
        if cid in chunks:
            raise ValueError(f"Duplicate corpus chunk_id: {cid}")
        chunks[cid] = {
            "chunk_id": cid,
            "tenant_id": _id(_field(row, "tenant_id")),
            "document_id": _id(_field(row, "document_id", "source_document_id")),
            "passage": str(_field(row, "passage", "text", "positive_text")),
        }
    queries = {}
    for row in cases:
        qid = _id(_field(row, "query_id"))
        if qid in queries:
            raise ValueError(f"Duplicate query_id: {qid}")
        gold = row.get("positive_chunk_ids")
        if gold is None:
            gold = row.get("positive_chunk_id", row.get("chunk_id"))
        if gold is None:
            raise ValueError(f"{qid}: no positive IDs")
        if not isinstance(gold, list):
            gold = [gold]
        gold = list(dict.fromkeys(_id(x) for x in gold))
        if not gold:
            raise ValueError(f"{qid}: empty gold IDs")
        tenant = _id(_field(row, "tenant_id"))
        missing = [cid for cid in gold if cid not in chunks]
        if missing:
            raise ValueError(
                f"{qid}: gold chunk(s) absent from exported corpus: {missing}. "
                "Cannot score or infer gold document identity."
            )
        for cid in gold:
            if chunks[cid]["tenant_id"] != tenant:
                raise ValueError(f"{qid}: gold chunk {cid} belongs to different tenant")
        queries[qid] = {
            "query_id": qid,
            "query": str(_field(row, "query_text", "query")),
            "tenant_id": tenant,
            "gold_chunk_ids": gold,
            "gold_document_ids": sorted({chunks[x]["document_id"] for x in gold}),
        }
    return queries, chunks

def _norm(value):
    return " ".join(str(value).casefold().split())


def _ids(row, plural, singular):
    values = row.get(plural)
    if values is None:
        values = row.get(singular)
    if values is None:
        return []
    if not isinstance(values, list):
        values = [values]
    return [_id(x) for x in values]


def check_split_overlap(cases, train_rows, chunks):
    """Conservatively inspect every known training passage and source document.

    Note: absence of collisions is NOT proof of independence if training examples
    cannot be mapped to source documents or if paraphrased queries are present.
    """
    from collections import defaultdict
    by_text = defaultdict(set)
    for chunk in chunks.values():
        by_text[_norm(chunk["passage"])].add(chunk["chunk_id"])

    train_queries, train_chunks, train_documents = set(), set(), set()
    untraceable = []
    for index, row in enumerate(train_rows):
        query = row.get("query_text") or row.get("query")
        if query:
            train_queries.add(_norm(query))
        ids = []
        for plural, singular in (
            ("positive_chunk_ids", "positive_chunk_id"),
            ("negative_chunk_ids", "negative_chunk_id"),
            ("candidate_chunk_ids", "candidate_chunk_id"),
        ):
            ids.extend(_ids(row, plural, singular))
        for field in ("chunk_id", "candidate_negative_chunk_id"):
            if row.get(field) is not None:
                ids.append(_id(row[field]))
        known_sources = set()
        for cid in ids:
            train_chunks.add(cid)
            if cid in chunks:
                known_sources.add(chunks[cid]["document_id"])
            else:
                untraceable.append({"row": index, "reason": f"unknown chunk ID: {cid}"})
        for field in ("document_id", "source_document_id",
                      "positive_document_id", "negative_document_id"):
            if row.get(field) is not None:
                known_sources.add(_id(row[field]))
        # Stage 2 spreadsheet exports often contain only raw passage text.
        # Resolve both positive AND negative against the exact exported corpus.
        for field in ("positive", "positive_text", "candidate_negative",
                      "negative", "negative_text"):
            passage = row.get(field)
            if not passage:
                continue
            matches = by_text.get(_norm(passage), set())
            if not matches:
                untraceable.append({"row": index, "reason": f"{field}: no exact passage match"})
            for cid in matches:
                train_chunks.add(cid)
                known_sources.add(chunks[cid]["document_id"])
        if not known_sources:
            untraceable.append({"row": index, "reason": "source document not established"})
        train_documents.update(known_sources)

    collisions = []
    for case in cases.values():
        found = []
        if _norm(case["query"]) in train_queries:
            found.append("query")
        if set(case["gold_chunk_ids"]) & train_chunks:
            found.append("gold_chunk")
        if set(case["gold_document_ids"]) & train_documents:
            found.append("source_document")
        if found:
            collisions.append({"query_id": case["query_id"], "collision_types": found})
    return collisions, untraceable

def validate_candidates(cases, chunks, rows, top_k=10):
    """Validate a frozen candidate file without silently dropping any query."""
    result = {}
    for row in rows:
        qid = _id(_field(row, "query_id"))
        if qid in result or qid not in cases:
            raise ValueError(f"Unexpected or duplicate candidates for query: {qid}")
        ids_raw = row.get("chunk_ids", row.get("candidate_chunk_ids"))
        if not isinstance(ids_raw, list) or not ids_raw:
            raise ValueError(f"{qid}: candidates must be a nonempty list")
        ids = [_id(x) for x in ids_raw]
        if len(ids) != len(set(ids)):
            raise ValueError(f"{qid}: duplicate candidate chunk")
        if len(ids) > top_k:
            raise ValueError(f"{qid}: more than {top_k} candidates")
        for cid in ids:
            if cid not in chunks:
                raise ValueError(f"{qid}: missing corpus chunk {cid}")
            if chunks[cid]["tenant_id"] != cases[qid]["tenant_id"]:
                raise ValueError(f"{qid}: cross-tenant candidate {cid}")
        result[qid] = ids
    if set(result) != set(cases):
        raise ValueError(
            f"Candidate query IDs differ from case IDs "
            f"(missing: {sorted(set(cases) - set(result))[:5]})"
        )
    return result

def build_qwen_candidates(cases, chunks, top_k=10,
                          model_id="Qwen/Qwen3-Embedding-0.6B"):
    """One research candidate set; filter by tenant before scoring/ranking."""
    import numpy as np
    from sentence_transformers import SentenceTransformer

    model = SentenceTransformer(model_id, device="cuda")
    model.max_seq_length = 256
    ids = list(chunks)
    passage_vectors = model.encode_document(
        [chunks[cid]["passage"] for cid in ids],
        normalize_embeddings=True, show_progress_bar=True,
    )
    case_ids = list(cases)
    query_vectors = model.encode_query(
        [cases[qid]["query"] for qid in case_ids],
        normalize_embeddings=True,
    )
    result = {}
    for qid, vector in zip(case_ids, query_vectors):
        allowed = [i for i, cid in enumerate(ids)
                   if chunks[cid]["tenant_id"] == cases[qid]["tenant_id"]]
        scores = np.asarray(passage_vectors[allowed]) @ np.asarray(vector)
        ranked = sorted(zip(allowed, scores.tolist()),
                        key=lambda pair: (-pair[1], ids[pair[0]]))
        result[qid] = [ids[i] for i, _ in ranked[:top_k]]
    return result


def rerank_fixed(query, candidate_ids, chunks, predict, pool_size=5):
    """Score every candidate, reorder only the head and leave the tail intact.

    All raw logits are diagnostic observations, including tail items that MUST
    NOT be promoted by the bounded serving-style top-N rerank.
    """
    pairs = [[query, chunks[cid]["passage"]] for cid in candidate_ids]
    raw_scores = list(predict(pairs)) if pairs else []
    if len(raw_scores) != len(candidate_ids):
        raise ValueError("Reranker returned wrong score count")
    by_chunk = {}
    for cid, value in zip(candidate_ids, raw_scores):
        try:
            score = float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid reranker score for {cid}") from exc
        if not math.isfinite(score):
            raise ValueError(f"Invalid reranker score for {cid}")
        by_chunk[cid] = score
    pool = candidate_ids[:pool_size]
    head = sorted(pool, key=lambda cid: -by_chunk[cid])  # stable score ties
    return head + candidate_ids[len(pool):], by_chunk

def _metrics(cases, rows):
    n = len(cases)
    def hit(k, key):
        return round(100 * sum(
            r[key] is not None and r[key] <= k for r in rows
        ) / n, 2)
    return {
        "queries": n,
        "hit_at_1": hit(1, "gold_rank"),
        "hit_at_5": hit(5, "gold_rank"),
        "hit_at_10": hit(10, "gold_rank"),
        "document_hit_at_1_unverified": hit(1, "gold_document_rank"),
        "document_hit_at_5_unverified": hit(5, "gold_document_rank"),
        "candidate_recall_at_10": hit(10, "original_gold_rank"),
        "mrr_at_10": round(sum(
            1 / r["gold_rank"] for r in rows
            if r["gold_rank"] is not None and r["gold_rank"] <= 10
        ) / n, 4),
    }

def evaluate_one(cases, chunks, candidates, predict, pool_size=5):
    rows = []
    for qid, case in cases.items():
        ordered, scores = rerank_fixed(
            case["query"], candidates[qid], chunks, predict, pool_size
        )
        gold = set(case["gold_chunk_ids"])
        documents = set(case["gold_document_ids"])
        rank = next((i for i, cid in enumerate(ordered, 1) if cid in gold), None)
        original_rank = next(
            (i for i, cid in enumerate(candidates[qid], 1) if cid in gold), None
        )
        doc_rank = next(
            (i for i, cid in enumerate(ordered, 1)
             if chunks[cid]["document_id"] in documents), None
        )
        rows.append({
            **case, "candidate_chunk_ids": candidates[qid],
            "ranked_chunk_ids": ordered, "reranker_scores": scores,
            "gold_rank": rank, "original_gold_rank": original_rank,
            "gold_document_rank": doc_rank,
        })
    return {"metrics": _metrics(cases, rows), "queries": rows}

def _sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as f:
        for part in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(part)
    return digest.hexdigest()


def checkpoint_fingerprint(path):
    """Hash local checkpoint files; remote model identifiers are not immutable."""
    p = Path(path)
    if not p.exists():
        if p.is_absolute() or path.startswith("."):
            raise FileNotFoundError(
                f"Local checkpoint missing: {path}. Do not retrain implicitly."
            )
        return {"source": path, "sha256": None,
                "status": "REMOTE_MODEL_REVISION_UNPINNED"}
    if p.is_file():
        return {"source": str(p), "sha256": _sha256(p), "status": "LOCAL_FILE"}
    h = hashlib.sha256()
    paths = sorted(x for x in p.rglob("*") if x.is_file())
    if not paths:
        raise ValueError(f"No checkpoint files found in {p}")
    for file in paths:
        rel = file.relative_to(p).as_posix()
        # File name length-prefix ensures no ambiguity from concatenation.
        key = rel.encode("utf-8")
        h.update(len(key).to_bytes(8, "big"))
        h.update(key)
        h.update(bytes.fromhex(_sha256(file)))
    return {"source": str(p), "sha256": h.hexdigest(),
            "file_count": len(paths), "status": "LOCAL_DIRECTORY"}


def load_training_input(path):
    """Read exact-used train examples from JSONL or reviewer-exported CSV/XLSX."""
    p = Path(path)
    if p.suffix.lower() == ".jsonl":
        return load_jsonl(p)
    if p.suffix.lower() == ".csv":
        with p.open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))
    if p.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl required to read reviewer XLSX") from exc
        book = load_workbook(p, read_only=True, data_only=True)
        try:
            if "AI Triage - 50" not in book.sheetnames:
                raise ValueError("Expected the 'AI Triage - 50' review sheet")
            ws = book["AI Triage - 50"]
            iterator = ws.values
            columns = [str(v or "").strip() for v in next(iterator)]
            return [dict(zip(columns, values)) for values in iterator
                    if any(v is not None for v in values)]
        finally:
            book.close()
    raise ValueError(f"Unsupported training input: {path}")
 

def run(args):
    cases, chunks = normalize_inputs(load_jsonl(args.cases), load_jsonl(args.corpus))
    collision_details, unresolved_provenance = [], []
    for training_path in args.train or []:
        collisions, unknown = check_split_overlap(
            cases, load_training_input(training_path), chunks
        )
        collision_details.extend({"input": training_path, **x} for x in collisions)
        unresolved_provenance.extend({"input": training_path, **x} for x in unknown)
    if collision_details:
        raise ValueError(
            f"TRAIN/CASE LEAKAGE: {json.dumps(collision_details[:20])}"
        )
    if unresolved_provenance and not args.allow_untraceable_training:
        raise ValueError(
            "UNTRACEABLE TRAINING SOURCES: cannot certify separation: "
            + json.dumps(unresolved_provenance[:20])
        )
    if args.candidates:
        candidates = validate_candidates(
            cases, chunks, load_jsonl(args.candidates), args.top_k,
        )
    else:
        candidates = build_qwen_candidates(
            cases, chunks, args.top_k, args.embedding_model,
        )
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    with (out / "frozen_candidates.jsonl").open("w", encoding="utf-8") as f:
        for qid, ids in candidates.items():
            f.write(json.dumps({"query_id": qid, "chunk_ids": ids}) + "\n")

    from sentence_transformers import CrossEncoder
    import torch
    results, fingerprints = {}, {}
    for spec in args.model:
        if "=" not in spec:
            raise ValueError(f"Invalid --model {spec!r}; expected label=path")
        name, path = spec.split("=", 1)
        if not name or name in results:
            raise ValueError("Duplicate or empty model label")
        fingerprints[name] = checkpoint_fingerprint(path)
        model = CrossEncoder(path, device=args.device, max_length=384)
        if model.num_labels != 1:
            raise ValueError(f"{name}: num_labels={model.num_labels}; expected one ranking logit")
        def predict(pairs):
            return model.predict(
                pairs, activation_fct=torch.nn.Identity(),
                apply_softmax=False, show_progress_bar=False,
            )
        results[name] = evaluate_one(
            cases, chunks, candidates, predict, args.pool_size,
        )
        del model
        if torch.cuda.is_available():
            torch.cuda.empty_cache()

    per_query = []
    for qid, case in cases.items():
        row = {
            "query_id": qid, "query": case["query"],
            "tenant_id": case["tenant_id"],
            "gold_chunk_ids": case["gold_chunk_ids"],
            "gold_document_ids": case["gold_document_ids"],
            "fixed_candidates": candidates[qid],
        }
        for name, report in results.items():
            query_result = next(r for r in report["queries"] if r["query_id"] == qid)
            row[name] = {
                "ranked_chunk_ids": query_result["ranked_chunk_ids"],
                "reranker_scores": query_result["reranker_scores"],
                "gold_rank": query_result["gold_rank"],
                "gold_document_rank": query_result["gold_document_rank"],
            }
        if results:
            baseline_name = next(iter(results))
            baseline_rank = row[baseline_name]["gold_rank"]
            for name in results:
                current_rank = row[name]["gold_rank"]
                row[name]["rank_change_vs_baseline"] = (
                    baseline_rank - current_rank
                    if baseline_rank is not None and current_rank is not None
                    else None
                )
        per_query.append(row)
    with (out / "per_query.jsonl").open("w", encoding="utf-8") as f:
        for row in per_query:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
    with (out / "per_query.csv").open("w", newline="", encoding="utf-8") as f:
        columns = ["query_id", "query", "tenant_id", "gold_chunk_ids",
                   "gold_document_ids", "fixed_candidates"]
        for name in results:
            columns += [f"{name}_ranked", f"{name}_scores",
                        f"{name}_gold_rank", f"{name}_document_rank",
                        f"{name}_rank_change_vs_baseline"]
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        for row in per_query:
            flat = {
                k: json.dumps(v) if isinstance(v, (list, dict)) else v
                for k, v in row.items() if k not in results
            }
            for name in results:
                data = row[name]
                flat[f"{name}_ranked"] = json.dumps(data["ranked_chunk_ids"])
                flat[f"{name}_scores"] = json.dumps(data["reranker_scores"])
                flat[f"{name}_gold_rank"] = data["gold_rank"]
                flat[f"{name}_document_rank"] = data["gold_document_rank"]
                flat[f"{name}_rank_change_vs_baseline"] = data["rank_change_vs_baseline"]
            writer.writerow(flat)

    summary = {
        "warning": "DEVELOPMENT DIAGNOSTIC; reused weak labels are not independent accuracy evidence.",
        "gold_document_note": "Inferred from corpus gold chunk metadata, not independently adjudicated.",
        "scoring": "Raw single-logit scores for all candidates; descending rank applied ONLY to the bounded head pool; tail retains Qwen order.",
        "cases_sha256": _sha256(args.cases),
        "corpus_sha256": _sha256(args.corpus),
        "candidates_sha256": _sha256(out / "frozen_candidates.jsonl"),
        "pool_size": args.pool_size, "top_k": args.top_k,
        "checkpoint_fingerprints": fingerprints,
        "training_inputs": list(args.train or []),
        "train_overlap_check": (
            "POTENTIALLY_INCOMPLETE" if unresolved_provenance else
            "PARTIAL_EXACT_PROVENANCE_CHECK" if args.train else
            "NOT_PERFORMED"
        ),
        "untraceable_training_sources": unresolved_provenance[:100],
        "holdout_independent": False,
        "production_authorization_tested": False,
        "results": {name: report["metrics"] for name, report in results.items()},
    }
    (out / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cases", required=True)
    parser.add_argument("--corpus", required=True)
    parser.add_argument("--train", action="append",
                        help="Exact training JSONL/CSV/XLSX; repeat for Stage 1 and exact Stage 2 rows")
    parser.add_argument("--allow-untraceable-training", action="store_true",
                        help="Research only: report unknown training provenance, do not treat splits as verified")
    parser.add_argument("--candidates", help="Optional frozen Qwen top-10 candidate JSONL")
    parser.add_argument("--model", action="append", required=True,
                        help="name=HF_model_or_local_checkpoint")
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--pool-size", type=int, default=5)
    parser.add_argument("--top-k", type=int, default=10)
    parser.add_argument("--device", default="cuda")
    parser.add_argument("--embedding-model", default="Qwen/Qwen3-Embedding-0.6B")
    args = parser.parse_args()
    if not 0 < args.pool_size <= args.top_k <= 25:
        parser.error("Require 0 < pool-size <= top-k <= 25")
    run(args)


if __name__ == "__main__":
    main()
